from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

doc = Document()

# ── Page margins ──
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# ── Default font ──
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
font.color.rgb = RGBColor(0, 0, 0)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

# ── Helper functions ──
def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_after = Pt(4)

def add_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(80, 80, 80)
    run.italic = True
    p.paragraph_format.space_after = Pt(14)

def add_heading_custom(text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Times New Roman'
    if level == 1:
        run.font.size = Pt(14)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(12)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)

def add_body(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        run_b = p.add_run(bold_prefix)
        run_b.bold = True
        run_b.font.name = 'Times New Roman'
        run_b.font.size = Pt(12)
        run_n = p.add_run(text)
        run_n.font.name = 'Times New Roman'
        run_n.font.size = Pt(12)
    else:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

# ════════════════════════════════════════════════════════
# DOCUMENT CONTENT
# ════════════════════════════════════════════════════════

add_title("Controllable Fashion Image Synthesis via\nLow-Rank Adaptation and Structural Guidance")
add_subtitle("Deep Learning Course — Project Proposal")

# ── 1. Introduction ──
add_heading_custom("1. Introduction and Motivation")
add_body(
    "The fashion industry relies heavily on visual prototyping and design iterations, "
    "a process traditionally requiring significant time, physical materials, and human "
    "resources to conceptualize new garments. The advent of Deep Learning, particularly "
    "Latent Diffusion Models (LDMs), has revolutionized image synthesis, allowing for "
    "the generation of highly realistic images from natural language descriptions. "
    "However, applying these generative models to professional fashion design presents "
    "unique challenges. Fashion synthesis demands not only high visual fidelity and "
    "domain-specific stylistic accuracy but also strict preservation of underlying "
    "structural elements, such as the silhouette, pose, and geometry of the garment."
)
add_body(
    "This project proposes the development of a controllable fashion image synthesis "
    "system. By leveraging advanced generative networks alongside structural condition "
    "mechanisms, this proposed framework aims to allow users to generate novel fashion "
    "items from text descriptions while strictly adhering to the spatial and structural "
    "constraints of a reference image."
)

# ── 2. Problem Statement ──
add_heading_custom("2. Problem Statement")
add_body(
    "While vanilla text-to-image models (such as DALL-E or standard Stable Diffusion) "
    "can generate visually appealing clothing, they lack fine-grained structural "
    "controllability. If a designer wishes to change the fabric of a jacket to \"black "
    "leather\" while maintaining the exact cut, collar shape, and draping of the original "
    "sketch, standard models typically alter the entire geometry of the image. Furthermore, "
    "broad-domain generative models often lack the specialized nuance required for "
    "high-end fashion textures and details."
)
add_body(
    "Therefore, the core problem this project addresses is: How can we bridge the gap "
    "between open-ended text-to-image generation and strict structural preservation "
    "within the specialized domain of fashion imagery?"
)

# ── 3. Proposed Methodology ──
add_heading_custom("3. Proposed Methodology")
add_body(
    "To solve the aforementioned challenges, we propose a multi-stage deep learning "
    "architecture that integrates domain fine-tuning with spatial conditioning. The "
    "proposed pipeline consists of three core components:"
)

add_heading_custom("3.1. Base Generative Model – Stable Diffusion v1.5", level=2)
add_body(
    "We will employ Stable Diffusion v1.5 as our foundational text-to-image framework. "
    "As a latent diffusion model, it offers a strong baseline for high-fidelity image "
    "synthesis. Operating in the latent space allows for significant computational "
    "efficiency compared to standard pixel-space diffusion models, making it feasible "
    "to train and evaluate within the constraints of academic computing resources."
)

add_heading_custom("3.2. Structural Guidance via ControlNet", level=2)
add_body(
    "To achieve structural constraints, we propose integrating a ControlNet module. "
    "ControlNet is a neural network architecture designed to add spatial conditioning "
    "controls to large, pre-trained text-to-image diffusion models. Before feeding a "
    "reference image into the network, we will apply a Canny edge detector to extract "
    "a structural map (outlining the silhouette, seams, and pose). ControlNet will take "
    "this structural map as an auxiliary input condition, locking the geometry of the "
    "generated output to the reference image, ensuring the generated garment retains "
    "the exact shape of the original."
)

add_heading_custom("3.3. Domain Adaptation via LoRA (Low-Rank Adaptation)", level=2)
add_body(
    "Because Stable Diffusion is trained on a general dataset, it requires adaptation "
    "to understand the nuanced textures, lighting, and semantic terminology of the "
    "fashion domain. Full fine-tuning of the model is computationally prohibitive. "
    "Instead, we propose using LoRA (Low-Rank Adaptation). LoRA injects trainable "
    "rank decomposition matrices into the transformer layers of the diffusion model "
    "while keeping the original pre-trained weights frozen. This allows us to achieve "
    "domain-specific fine-tuning by updating only a fraction (~4 million) of the "
    "parameters versus the full 860 million, making the training process highly efficient."
)

# ── 4. Dataset ──
add_heading_custom("4. Dataset")
add_body(
    "We plan to utilize the FashionGen dataset for our fine-tuning and evaluation phases. "
    "The dataset provides high-resolution images of fashion models and isolated garments, "
    "paired with detailed, professional textual descriptions."
)
add_bullet(" The dataset will be pre-processed to extract Canny edge maps for every image.")
add_bullet(
    " We will configure a training split (30,000 to 100,000 images depending on hardware "
    "limitations) and a hold-out evaluation split to measure generative accuracy and "
    "generalization."
)

# ── 5. Expected Outcomes ──
add_heading_custom("5. Expected Outcomes")
add_body("Based on the proposed methodology, we expect to successfully build a system capable of:")
add_bullet(" Generating highly realistic garments that accurately reflect complex textual prompts.", bold_prefix="Accurate Text-to-Image Translation:")
add_bullet(" Demonstrating that the generated images perfectly align with the edges and pose of the provided reference image.", bold_prefix="Structural Preservation:")
add_bullet(" Proving that LoRA allows for high-quality domain adaptation without needing industrial-scale compute clusters.", bold_prefix="Resource Efficiency:")
add_body(
    "Ultimately, the expected outcome is a proof-of-concept pipeline that can seamlessly "
    "transition a basic garment sketch or reference photo into various high-quality "
    "stylistic renders based entirely on text."
)

# ── 6. Evaluation Strategy ──
add_heading_custom("6. Evaluation Strategy")
add_body(
    "Because generative tasks are notoriously difficult to evaluate objectively, we will "
    "not rely solely on human visual inspection. Once the model is built, we propose "
    "utilizing a rigorous suite of quantitative metrics:"
)
add_bullet(" To measure the realism and distribution similarity between our generated fashion images and the real dataset.", bold_prefix="Fréchet Inception Distance (FID) & Kernel Inception Distance (KID):")
add_bullet(" To evaluate the perceptual similarity and structural retention between the output and the original reference.", bold_prefix="LPIPS (Learned Perceptual Image Patch Similarity):")
add_bullet(" To measure the semantic alignment between the user's text prompt and the generated image.", bold_prefix="CLIP Score:")

# ── 7. Timeline ──
add_heading_custom("7. Proposed Timeline")
add_bullet(" Dataset acquisition, pre-processing, and extraction of Canny edge maps.", bold_prefix="Week 1–2:")
add_bullet(" Environment setup, deployment of Stable Diffusion, and integration of ControlNet.", bold_prefix="Week 3–4:")
add_bullet(" Training of the LoRA weights on the fashion dataset. Hyper-parameter tuning.", bold_prefix="Week 5–6:")
add_bullet(" Evaluation dataset generation, metric calculations (FID, LPIPS, CLIP), and analysis.", bold_prefix="Week 7–8:")
add_bullet(" Finalizing the report, creating visual comparison charts, and preparing the presentation.", bold_prefix="Week 9–10:")

# ── Save ──
output_path = r"c:\Users\muham\Documents\DL Project\Project_Proposal.docx"
doc.save(output_path)
print(f"Document saved to: {output_path}")
